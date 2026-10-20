import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Styles helper
HEADER_FILL = PatternFill(start_color="1B5E20", end_color="1B5E20", fill_type="solid") # Deep Forest Green Koperasi
SUBHEADER_FILL = PatternFill(start_color="E8F5E9", end_color="E8F5E9", fill_type="solid")
HEADER_FONT = Font(name="Arial", size=11, bold=True, color="FFFFFF")
TITLE_FONT = Font(name="Arial", size=14, bold=True, color="1B5E20")
SUBTITLE_FONT = Font(name="Arial", size=10, italic=True, color="555555")
BOLD_FONT = Font(name="Arial", size=10, bold=True)
REGULAR_FONT = Font(name="Arial", size=10)

THIN_BORDER = Border(
    left=Side(style='thin', color='CCCCCC'),
    right=Side(style='thin', color='CCCCCC'),
    top=Side(style='thin', color='CCCCCC'),
    bottom=Side(style='thin', color='CCCCCC')
)

TOTAL_BORDER = Border(
    left=Side(style='thin', color='333333'),
    right=Side(style='thin', color='333333'),
    top=Side(style='thin', color='333333'),
    bottom=Side(style='double', color='333333')
)

def auto_fit_columns(ws, max_cols=15):
    for col in range(1, max_cols + 1):
        col_letter = get_column_letter(col)
        max_len = 0
        for cell in ws[col_letter]:
            if cell.value:
                # simple length estimation
                val_str = str(cell.value)
                max_len = max(max_len, len(val_str))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

def export_rekap_simpanan_excel(members):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Rekap Simpanan"

    # Title header
    ws.merge_cells("A1:H1")
    ws["A1"] = "KOSITAS - KOPERASI STATISTIK TASIKMALAYA"
    ws["A1"].font = TITLE_FONT
    ws["A1"].alignment = Alignment(horizontal="center")

    ws.merge_cells("A2:H2")
    ws["A2"] = "REKAPITULASI SIMPANAN ANGGOTA"
    ws["A2"].font = Font(name="Arial", size=12, bold=True)
    ws["A2"].alignment = Alignment(horizontal="center")

    ws.merge_cells("A3:H3")
    ws["A3"] = "BPS Kabupaten Tasikmalaya - Jl. Raya Timur Singaparna KM. 4 Cintaraja"
    ws["A3"].font = SUBTITLE_FONT
    ws["A3"].alignment = Alignment(horizontal="center")

    headers = [
        "No", "No Anggota", "Nama Anggota", "Simpanan Pokok",
        "Simpanan Wajib", "Simpanan Lainnya", "Simpanan Lebaran", "Simpanan Qurban", "Jumlah Simpanan"
    ]

    row_num = 5
    for col_num, h in enumerate(headers, 1):
        cell = ws.cell(row=row_num, column=col_num, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    row_num = 6
    tot_pokok = tot_wajib = tot_manasuka = tot_lebaran = tot_qurban = tot_all = 0

    for idx, m in enumerate(members, 1):
        ws.cell(row=row_num, column=1, value=idx).alignment = Alignment(horizontal="center")
        ws.cell(row=row_num, column=2, value=m['no_anggota']).alignment = Alignment(horizontal="center")
        ws.cell(row=row_num, column=3, value=m['nama'])
        
        # Values
        c_pokok = ws.cell(row=row_num, column=4, value=m['simpanan_pokok'])
        c_wajib = ws.cell(row=row_num, column=5, value=m['simpanan_wajib'])
        c_man = ws.cell(row=row_num, column=6, value=m['simpanan_manasuka'])
        c_leb = ws.cell(row=row_num, column=7, value=m['simpanan_lebaran'])
        c_qur = ws.cell(row=row_num, column=8, value=m['simpanan_qurban'])
        c_tot = ws.cell(row=row_num, column=9, value=m['total_simpanan'])

        for c in [c_pokok, c_wajib, c_man, c_leb, c_qur, c_tot]:
            c.number_format = '#,##0'
            c.alignment = Alignment(horizontal="right")

        for col in range(1, 10):
            ws.cell(row=row_num, column=col).border = THIN_BORDER
            ws.cell(row=row_num, column=col).font = REGULAR_FONT

        tot_pokok += m['simpanan_pokok']
        tot_wajib += m['simpanan_wajib']
        tot_manasuka += m['simpanan_manasuka']
        tot_lebaran += m['simpanan_lebaran']
        tot_qurban += m['simpanan_qurban']
        tot_all += m['total_simpanan']
        row_num += 1

    # Total row
    ws.merge_cells(start_row=row_num, start_column=1, end_row=row_num, end_column=3)
    tot_cell = ws.cell(row=row_num, column=1, value="TOTAL KESELURUHAN")
    tot_cell.font = BOLD_FONT
    tot_cell.alignment = Alignment(horizontal="center")

    totals = [tot_pokok, tot_wajib, tot_manasuka, tot_lebaran, tot_qurban, tot_all]
    for i, t_val in enumerate(totals, 4):
        c = ws.cell(row=row_num, column=i, value=t_val)
        c.font = BOLD_FONT
        c.number_format = '#,##0'
        c.alignment = Alignment(horizontal="right")

    for col in range(1, 10):
        cell = ws.cell(row=row_num, column=col)
        cell.border = TOTAL_BORDER
        cell.fill = SUBHEADER_FILL

    auto_fit_columns(ws, 9)

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out

def export_data_piutang_excel(loans_data):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data Piutang"

    # Title header
    ws.merge_cells("A1:G1")
    ws["A1"] = "KOSITAS - KOPERASI STATISTIK TASIKMALAYA"
    ws["A1"].font = TITLE_FONT
    ws["A1"].alignment = Alignment(horizontal="center")

    ws.merge_cells("A2:G2")
    ws["A2"] = "DATA PIUTANG & PINJAMAN ANGGOTA"
    ws["A2"].font = Font(name="Arial", size=12, bold=True)
    ws["A2"].alignment = Alignment(horizontal="center")

    headers = [
        "No", "No Anggota", "Nama Anggota", "Pinjaman Japan",
        "Pinjaman Japen", "Pinjaman Bon Putih", "Pinjaman Barang", "Jumlah Piutang"
    ]

    row_num = 4
    for col_num, h in enumerate(headers, 1):
        cell = ws.cell(row=row_num, column=col_num, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")

    row_num = 5
    tot_japan = tot_japen = tot_bon = tot_brg = tot_all = 0

    for idx, m in enumerate(loans_data, 1):
        if m['total_piutang'] == 0:
            continue
        ws.cell(row=row_num, column=1, value=idx).alignment = Alignment(horizontal="center")
        ws.cell(row=row_num, column=2, value=m['no_anggota']).alignment = Alignment(horizontal="center")
        ws.cell(row=row_num, column=3, value=m['nama'])

        c_jpn = ws.cell(row=row_num, column=4, value=m['piutang_japan'])
        c_jpnk = ws.cell(row=row_num, column=5, value=m['piutang_japen'])
        c_bon = ws.cell(row=row_num, column=6, value=m['piutang_bon_putih'])
        c_brg = ws.cell(row=row_num, column=7, value=m['piutang_barang'])
        c_tot = ws.cell(row=row_num, column=8, value=m['total_piutang'])

        for c in [c_jpn, c_jpnk, c_bon, c_brg, c_tot]:
            c.number_format = '#,##0'
            c.alignment = Alignment(horizontal="right")

        for col in range(1, 9):
            ws.cell(row=row_num, column=col).border = THIN_BORDER
            ws.cell(row=row_num, column=col).font = REGULAR_FONT

        tot_japan += m['piutang_japan']
        tot_japen += m['piutang_japen']
        tot_bon += m['piutang_bon_putih']
        tot_brg += m['piutang_barang']
        tot_all += m['total_piutang']
        row_num += 1

    # Total row
    ws.merge_cells(start_row=row_num, start_column=1, end_row=row_num, end_column=3)
    tot_cell = ws.cell(row=row_num, column=1, value="TOTAL PIUTANG")
    tot_cell.font = BOLD_FONT
    tot_cell.alignment = Alignment(horizontal="center")

    totals = [tot_japan, tot_japen, tot_bon, tot_brg, tot_all]
    for i, t_val in enumerate(totals, 4):
        c = ws.cell(row=row_num, column=i, value=t_val)
        c.font = BOLD_FONT
        c.number_format = '#,##0'
        c.alignment = Alignment(horizontal="right")

    for col in range(1, 9):
        cell = ws.cell(row=row_num, column=col)
        cell.border = TOTAL_BORDER
        cell.fill = SUBHEADER_FILL

    auto_fit_columns(ws, 8)

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out

def export_rekap_bulanan_excel(rekap_data):
    wb = openpyxl.Workbook()
    ws = wb.active
    bulan_buku = rekap_data['bulan_buku']
    ws.title = f"Rekap {bulan_buku}"

    # Header
    ws.merge_cells("A1:K1")
    ws["A1"] = "KOSITAS - KOPERASI STATISTIK TASIKMALAYA"
    ws["A1"].font = TITLE_FONT
    ws["A1"].alignment = Alignment(horizontal="center")

    ws.merge_cells("A2:K2")
    ws["A2"] = f"REKAPITULASI PEMOTONGAN GAJI / SETORAN BULAN {bulan_buku}"
    ws["A2"].font = Font(name="Arial", size=12, bold=True)
    ws["A2"].alignment = Alignment(horizontal="center")

    headers = [
        "No", "No Anggota", "Nama Anggota", "Simpanan Wajib",
        "Simpanan Lebaran", "Simpanan Qurban", "Simpanan Manasuka",
        "Total Simpanan", "Angsuran Pokok", "Margin/Jasa", "Total Potongan"
    ]

    row_num = 4
    for col_num, h in enumerate(headers, 1):
        cell = ws.cell(row=row_num, column=col_num, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")

    row_num = 5
    for idx, item in enumerate(rekap_data.get('anggota_list', rekap_data.get('items', [])), 1):
        ws.cell(row=row_num, column=1, value=idx).alignment = Alignment(horizontal="center")
        ws.cell(row=row_num, column=2, value=item['no_anggota']).alignment = Alignment(horizontal="center")
        ws.cell(row=row_num, column=3, value=item['nama'])

        vals = [
            item['simpanan_wajib'],
            item['simpanan_lebaran'],
            item['simpanan_qurban'],
            item['simpanan_manasuka'],
            item['total_simpanan'],
            item['angsuran_pokok'],
            item['angsuran_margin'],
            item['total_potongan']
        ]

        for i, val in enumerate(vals, 4):
            c = ws.cell(row=row_num, column=i, value=val)
            c.number_format = '#,##0'
            c.alignment = Alignment(horizontal="right")

        for col in range(1, 12):
            ws.cell(row=row_num, column=col).border = THIN_BORDER
            ws.cell(row=row_num, column=col).font = REGULAR_FONT

        row_num += 1

    # Total Row
    ws.merge_cells(start_row=row_num, start_column=1, end_row=row_num, end_column=3)
    tot_cell = ws.cell(row=row_num, column=1, value="TOTAL BULAN INI")
    tot_cell.font = BOLD_FONT
    tot_cell.alignment = Alignment(horizontal="center")

    summary_vals = [
        rekap_data['total_wajib'],
        rekap_data['total_lebaran'],
        rekap_data['total_qurban'],
        rekap_data['total_manasuka'],
        rekap_data['total_simpanan'],
        rekap_data['total_angsuran_pokok'],
        rekap_data['total_angsuran_margin'],
        rekap_data['total_pemotongan']
    ]

    for i, val in enumerate(summary_vals, 4):
        c = ws.cell(row=row_num, column=i, value=val)
        c.font = BOLD_FONT
        c.number_format = '#,##0'
        c.alignment = Alignment(horizontal="right")

    for col in range(1, 12):
        cell = ws.cell(row=row_num, column=col)
        cell.border = TOTAL_BORDER
        cell.fill = SUBHEADER_FILL

    auto_fit_columns(ws, 11)

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out

def export_rekap_tahunan_excel(rekap_data):
    wb = openpyxl.Workbook()
    ws = wb.active
    tahun = rekap_data['tahun']
    ws.title = f"Rekap {tahun}"

    # Header
    ws.merge_cells("A1:L1")
    ws["A1"] = "KOSITAS - KOPERASI STATISTIK TASIKMALAYA"
    ws["A1"].font = TITLE_FONT
    ws["A1"].alignment = Alignment(horizontal="center")

    ws.merge_cells("A2:L2")
    ws["A2"] = f"REKAPITULASI TRANSAKSI & PEMOTONGAN GAJI TAHUN BUKU {tahun}"
    ws["A2"].font = Font(name="Arial", size=12, bold=True)
    ws["A2"].alignment = Alignment(horizontal="center")

    headers = [
        "No", "No Anggota", "Nama Anggota", "Simpanan Wajib",
        "Simpanan Lebaran", "Simpanan Qurban", "Simpanan Manasuka",
        "Total Simpanan", "Angsuran Pokok", "Margin/Jasa", "Total Potongan", "Catatan Pelunasan/TopUp"
    ]

    row_num = 4
    for col_num, h in enumerate(headers, 1):
        cell = ws.cell(row=row_num, column=col_num, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")

    row_num = 5
    for idx, item in enumerate(rekap_data.get('anggota_list', []), 1):
        ws.cell(row=row_num, column=1, value=idx).alignment = Alignment(horizontal="center")
        ws.cell(row=row_num, column=2, value=item['no_anggota']).alignment = Alignment(horizontal="center")
        ws.cell(row=row_num, column=3, value=item['nama'])

        vals = [
            item['simpanan_wajib'],
            item['simpanan_lebaran'],
            item['simpanan_qurban'],
            item['simpanan_manasuka'],
            item['total_simpanan'],
            item['angsuran_pokok'],
            item['angsuran_margin'],
            item['total_potongan']
        ]

        for i, val in enumerate(vals, 4):
            c = ws.cell(row=row_num, column=i, value=val)
            c.number_format = '#,##0'
            c.alignment = Alignment(horizontal="right")

        ws.cell(row=row_num, column=12, value=item.get('catatan_khusus', ''))

        for col in range(1, 13):
            ws.cell(row=row_num, column=col).border = THIN_BORDER
            ws.cell(row=row_num, column=col).font = REGULAR_FONT

        row_num += 1

    # Total Row
    ws.merge_cells(start_row=row_num, start_column=1, end_row=row_num, end_column=3)
    tot_cell = ws.cell(row=row_num, column=1, value=f"TOTAL TAHUN {tahun}")
    tot_cell.font = BOLD_FONT
    tot_cell.alignment = Alignment(horizontal="center")

    summary_vals = [
        rekap_data['total_wajib'],
        rekap_data['total_lebaran'],
        rekap_data['total_qurban'],
        rekap_data['total_manasuka'],
        rekap_data['total_simpanan'],
        rekap_data['total_angsuran_pokok'],
        rekap_data['total_angsuran_margin'],
        rekap_data['total_pemotongan']
    ]

    for i, val in enumerate(summary_vals, 4):
        c = ws.cell(row=row_num, column=i, value=val)
        c.font = BOLD_FONT
        c.number_format = '#,##0'
        c.alignment = Alignment(horizontal="right")

    for col in range(1, 13):
        cell = ws.cell(row=row_num, column=col)
        cell.border = TOTAL_BORDER
        cell.fill = SUBHEADER_FILL

    auto_fit_columns(ws, 12)

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out

def export_shu_excel(shu_data):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Daftar Pembagian SHU"

    # Header
    ws.merge_cells("A1:G1")
    ws["A1"] = "KOSITAS - KOPERASI STATISTIK TASIKMALAYA"
    ws["A1"].font = TITLE_FONT
    ws["A1"].alignment = Alignment(horizontal="center")

    ws.merge_cells("A2:G2")
    ws["A2"] = "DAFTAR PEMBAGIAN SISA HASIL USAHA (SHU) ANGGOTA"
    ws["A2"].font = Font(name="Arial", size=12, bold=True)
    ws["A2"].alignment = Alignment(horizontal="center")

    headers = [
        "No", "No Anggota", "Nama Anggota", "Jumlah Simpanan",
        "Jumlah Jasa Pinjaman", "SHU Jasa Modal", "SHU Jasa Pinjaman", "SHU Diterima"
    ]

    row_num = 4
    for col_num, h in enumerate(headers, 1):
        cell = ws.cell(row=row_num, column=col_num, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")

    row_num = 5
    tot_simp = tot_jasa = tot_shu_m = tot_shu_p = tot_shu_all = 0

    for idx, r in enumerate(shu_data['rincian_anggota'], 1):
        ws.cell(row=row_num, column=1, value=idx).alignment = Alignment(horizontal="center")
        ws.cell(row=row_num, column=2, value=r['no_anggota']).alignment = Alignment(horizontal="center")
        ws.cell(row=row_num, column=3, value=r['nama'])

        c_s = ws.cell(row=row_num, column=4, value=r['total_simpanan'])
        c_j = ws.cell(row=row_num, column=5, value=r['total_jasa_pinjaman'])
        c_sm = ws.cell(row=row_num, column=6, value=r['shu_jasa_modal'])
        c_sp = ws.cell(row=row_num, column=7, value=r['shu_jasa_pinjaman'])
        c_st = ws.cell(row=row_num, column=8, value=r['shu_total_diterima'])

        for c in [c_s, c_j, c_sm, c_sp, c_st]:
            c.number_format = '#,##0'
            c.alignment = Alignment(horizontal="right")

        for col in range(1, 9):
            ws.cell(row=row_num, column=col).border = THIN_BORDER
            ws.cell(row=row_num, column=col).font = REGULAR_FONT

        tot_simp += r['total_simpanan']
        tot_jasa += r['total_jasa_pinjaman']
        tot_shu_m += r['shu_jasa_modal']
        tot_shu_p += r['shu_jasa_pinjaman']
        tot_shu_all += r['shu_total_diterima']
        row_num += 1

    # Total Row
    ws.merge_cells(start_row=row_num, start_column=1, end_row=row_num, end_column=3)
    tot_cell = ws.cell(row=row_num, column=1, value="TOTAL SHU ANGGOTA")
    tot_cell.font = BOLD_FONT
    tot_cell.alignment = Alignment(horizontal="center")

    totals = [tot_simp, tot_jasa, tot_shu_m, tot_shu_p, tot_shu_all]
    for i, val in enumerate(totals, 4):
        c = ws.cell(row=row_num, column=i, value=val)
        c.font = BOLD_FONT
        c.number_format = '#,##0'
        c.alignment = Alignment(horizontal="right")

    for col in range(1, 9):
        cell = ws.cell(row=row_num, column=col)
        cell.border = TOTAL_BORDER
        cell.fill = SUBHEADER_FILL

    auto_fit_columns(ws, 8)

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out
